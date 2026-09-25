"""Export worker: JSON-lines protocol over stdin/stdout (spec section 6).

One conversion job at a time. stdout carries protocol messages only; stderr
carries bounded, redacted diagnostics (never document text or paths).
"""
from __future__ import annotations

import hashlib
import json
import os
import queue
import sys
import threading
import zipfile
from importlib import metadata
from pathlib import Path
from typing import Any

from . import BACKEND_ID, PROTOCOL_VERSION
from .errors import ExportError
from .reconstruct import ExportOptions, build_document
from .writers.docx_writer import write_docx

_ARTIFACT_NAMES = {"docx": "document.docx", "xlsx": "document.xlsx", "html": "document.html", "md": "document.md",
                   "pptx": "document.pptx", "rtf": "document.rtf", "xml": "document.xml", "epub": "document.epub"}
# options each format accepts (a request naming another is refused, not ignored)
_FORMAT_OPTIONS = {
    "docx": {"include_images", "include_hyperlinks", "include_form_values", "include_comments",
             "running_headers", "layout_mode", "ocr", "locale", "allow_partial"},
    "xlsx": {"include_images", "include_hyperlinks", "include_form_values", "include_comments",
             "running_headers", "ocr", "locale", "allow_partial"},
    "html": {"include_images", "include_hyperlinks", "include_form_values", "include_comments",
             "running_headers", "layout_mode", "ocr", "locale", "allow_partial"},
    "md": {"include_images", "include_hyperlinks", "include_form_values", "include_comments",
           "running_headers", "ocr", "locale", "allow_partial", "markdown_images", "markdown_asset_folder"},
    "pptx": {"include_images", "include_hyperlinks", "include_form_values", "include_comments",
             "running_headers", "ocr", "locale", "allow_partial", "pptx_mode"},
    "rtf": {"include_images", "include_hyperlinks", "include_form_values", "include_comments",
            "running_headers", "ocr", "locale", "allow_partial"},
    "xml": {"include_images", "include_hyperlinks", "include_form_values", "include_comments",
            "running_headers", "ocr", "locale", "allow_partial", "xml_images"},
    "epub": {"include_images", "include_hyperlinks", "include_form_values", "include_comments",
             "running_headers", "ocr", "locale", "allow_partial"},
}
EPUB_FIDELITY_LIMITS = [
    "a reflowable EPUB 3 book in reading order, built from Responsive reading: headings, paragraphs, lists, tables, figures, links, form values, comments; one chapter per source page, opening with a page break named after it",
    "page layout, fonts, colours, rules and shading are not reproduced (EPUB_LAYOUT_NOT_KEPT); the reading system sets the type",
    "the table of contents is the PDF's bookmarks that point into the selected pages; without bookmarks it is the headings; a page list maps every source page",
    "pictures are files in the package at no more than 200 dpi for their placed size (IMAGE_DOWNSAMPLED), with the labels they keep as alt text",
    "a scan's OCR layer is in a collapsed section after its page; it may contain recognition errors",
    "language is 'und' (none is known); the title is the PDF's, when it has one",
    "no scripts and no remote resources",
]
XML_FIDELITY_LIMITS = [
    "structured data, not a rendition: pages; in reading order, text blocks (paragraphs, headings, list items, invisible text layers) with their lines and words (box, PDF font name, size, bold, italic, link), tables (cells with spans, header flags, text, words), images, form fields and comments; schema zpdf-document-1.xsd (docs/xml-format.md), validated on every result",
    "coordinates are points on the page as displayed (after /Rotate), origin top-left, y down",
    "reading order is the reconstruction's: every item has a document-wide order number; words are grouped into lines by geometry",
    "rules, shading, colours and vector drawing are not described; a drawing that is rendered is an image (origin='vector')",
    "images are base64 PNG or JPEG at no more than 200 dpi for their placed size (xml_images='embed', default) or described without their pixels (xml_images='omit', XML_IMAGE_DATA_OMITTED)",
    "text inside check-box sentences has its text but no word boxes",
    "characters XML 1.0 cannot hold (control characters) become U+FFFD (XML_CHARS_REPLACED)",
]
RTF_FIDELITY_LIMITS = [
    "reading order, not page position: headings (outline levels), paragraphs in their dominant font and size, bold, italic, links, lists, tables, pictures, form values and comments",
    "never kept (RTF_LAYOUT_NOT_KEPT, on every result): page layout and positions, columns, rules and shading, text colour",
    "fonts are the installed stand-ins (Arial, Times New Roman, Courier New) named from the PDF fonts; a paragraph takes its dominant font and size, so a change of font inside a paragraph is not kept",
    "lists keep the source's numbering style and start (decimal, letters, roman) and bullets; the marker is also written as the list's fallback text",
    "tables have header rows repeated, column widths from the source (scaled to 6.5 in) and cells that span columns or rows; cell text is 8 pt Arial",
    "pictures are embedded as PNG or JPEG at no more than 200 dpi for their placed size (IMAGE_DOWNSAMPLED), with the labels they keep as their description",
    "comments are paragraphs in a 'Comments (page N)' section, not annotations (RTF_COMMENTS_AS_TEXT)",
    "invisible text (a scan's OCR layer) is hidden text",
    "the file is 7-bit RTF; it contains no field other than HYPERLINK (http, https, mailto) and no embedded object",
]
PPTX_FIDELITY_LIMITS = {
    "editable": [
        "one slide per page; a presentation has one slide size: the largest of the pages within 10 % of the most common size, those pages unscaled at the top left (PPTX_PAGE_PADDED); a page further off is scaled to fit and centred, fonts included (PPTX_PAGE_SCALED)",
        "every paragraph is one text box whose paragraphs are the source's lines at their measured positions; lines are not re-wrapped, so edited text does not reflow across lines",
        "text is set in an installed stand-in font (Arial, Times New Roman or Courier New), fitted to each line's measured width with character spacing; glyph shapes differ from the PDF's embedded fonts",
        "filled rectangles and rules are rectangle shapes; curves, gradients and other vector artwork are pictures, not editable shapes",
        "tables are drawn (rules and shading as shapes, one text box per cell), not PowerPoint tables",
        "a page displayed rotated is one group turned by the display rotation",
        "a scan's OCR layer is transparent, selectable text over the scan picture",
        "comments are written to the slide's speaker notes (PPTX_COMMENTS_AS_NOTES)",
        "pictures are embedded at no more than 200 dpi for their placed size (IMAGE_DOWNSAMPLED)",
    ],
    "page_image": [
        "the page's drawing (fills, rules, paths, shading) without its text and pictures is one background picture per slide, rendered at 200 dpi; it cannot be edited as shapes",
        "pictures and text are placed over it as in editable mode",
        "text inside Form XObjects cannot be removed from the background and would show twice",
    ],
}
MD_FIDELITY_LIMITS = [
    "GitHub Flavored Markdown in reading order: headings, paragraphs, bold, italic, strikethrough, links, lists, pipe tables, pictures, form values, comments",
    "never kept (MD_FORMATTING_NOT_KEPT, on every result): page layout, fonts and sizes, colours, underline and highlighting, alignment, rules and shading",
    "pipe tables have one header row and no spans: multi-row headers are joined per column (MD_TABLE_HEADER_FLATTENED); spanning cells are written in their first cell (MD_TABLE_SPANS_NOT_KEPT)",
    "lettered or roman list markers are written as digits (MD_LIST_MARKERS_AS_DIGITS)",
    "pictures are files in a sibling folder, linked relatively (markdown_images='folder', default), or data URIs in the one file (markdown_images='embed'; some renderers, GitHub among them, do not display them: MD_PICTURES_AS_DATA_URI); include_images=false omits them",
    "comments are block quotes in reading order, not attached to the text they mark (MD_COMMENTS_AS_QUOTES)",
    "check boxes are task-list items; values inside sentences stay as check-box glyphs",
    "invisible text (a scan's OCR layer, the text layer of a graphic) is in a collapsed <details> section after its page; it may contain recognition errors",
]
HTML_FIDELITY_LIMITS = {
    "preserve": [
        "one fixed-size box per page; text lines, filled rectangles, rules and pictures at their measured positions in the source's paint order",
        "text is set in an installed stand-in font (Arial, Times New Roman or Courier New) scaled to each line's measured width; glyph shapes differ from the PDF's embedded fonts",
        "tables are drawn (rules, shading and cell text in place), not marked up as HTML tables; use Responsive reading for table structure",
        "curves, gradients and other vector artwork are pictures, not editable shapes",
        "a scan's OCR layer is transparent, selectable text over the scan picture",
        "pictures are embedded at no more than 200 dpi for their placed size (IMAGE_DOWNSAMPLED)",
        "comments appear as outlined areas whose text shows on hover or focus",
        "no scripts and no external resources; the file is self-contained",
    ],
    "reflow": [
        "reading order, not page position: headings, paragraphs, lists, tables, figures, links, form values and comments in one column that reflows from phone to desktop width",
        "page geometry, fonts, colours, rules and shading are not reproduced; bold, italic and links are",
        "tables carry header rows and merged cells; unruled bodies under a ruled header are recovered by column alignment (TABLE_BODY_RECOVERED)",
        "list markers are the browser's (numbered lists keep their start number); a source bullet glyph is not kept",
        "invisible text (a scan's OCR layer, the text layer of a graphic) is in a collapsed section after its page; it may contain recognition errors",
        "pictures are embedded at no more than 200 dpi for their placed size (IMAGE_DOWNSAMPLED)",
        "no scripts and no external resources; the file is self-contained",
    ],
}
XLSX_FIDELITY_LIMITS = [
    "one worksheet per reconstructed table; tables with identical headers that continue across page halves or pages are joined, the repeated header written once",
    "a cell is a number only when the number, in its number format, displays exactly as printed; leading-zero identifiers, dates, parenthesised labels, footnote-marked values and ambiguous text stay text",
    "no formulas are written; printed totals are values",
    "a true minus sign (U+2212) is stored as a negative number and shown with an ASCII hyphen",
    "unruled table bodies under a ruled header are read by column alignment (TABLE_BODY_RECOVERED); tables with no ruled header at all are not detected and stay text",
    "dot leaders are dropped; footnote marks are kept as superscript characters",
    "text outside tables goes to a Text sheet in reading order, invisible text layers as 'invisible text' rows; pictures are not carried (a placeholder row and XLSX_PICTURE_OMITTED)",
    "cell fonts, colours, borders and shading are not reproduced; header rows are bold and frozen",
    "a cell carries at most one hyperlink",
]
_KNOWN_REQUEST_KEYS = {"protocol_version", "operation", "job_id", "snapshot", "staging_directory",
                       "format", "pages", "options"}
_KNOWN_OPTIONS = {"include_images", "include_hyperlinks", "include_form_values", "include_comments",
                  "running_headers", "layout_mode", "ocr", "locale", "allow_partial",
                  "markdown_images", "markdown_asset_folder", "pptx_mode", "xml_images"}
_ASSET_FOLDER = __import__("re").compile(r"^\w[\w ._()\-]{0,99}$")


def versions() -> dict[str, str]:
    import pypdfium2_raw
    vj = json.loads((Path(pypdfium2_raw.__file__).parent / "version.json").read_text())
    pdfium = f"{vj['major']}.{vj['minor']}.{vj['build']}.{vj['patch']}"
    return {
        "helper": BACKEND_ID,
        "python": ".".join(map(str, sys.version_info[:3])),
        "pypdfium2": metadata.version("pypdfium2"),
        "pdfium": pdfium,
        "python_docx": metadata.version("python-docx"),
        "lxml": metadata.version("lxml"),
        "openpyxl": metadata.version("openpyxl"),
        "python_pptx": metadata.version("python-pptx"),
    }


def capabilities() -> dict[str, Any]:
    return {
        "protocol_version": PROTOCOL_VERSION,
        "operation": "capabilities",
        "backend": BACKEND_ID,
        "versions": versions(),
        "formats": {
            "docx": {
                "available": True,
                "options": {"include_images": [True, False], "include_hyperlinks": [True, False],
                            "include_form_values": [True, False], "include_comments": [True, False],
                            "running_headers": ["preserve"], "layout_mode": ["preserve", "reflow"],
                            "ocr": ["off"], "locale": [None],
                            "allow_partial": [True, False]},
                "unsupported_options": {"ocr": "OCR is not implemented", "locale": "no locale rules implemented",
                                        "running_headers=remove": "running-header removal not implemented"},
                "verified_fixture_scope": ["single-column memo", "ruled numeric table",
                                           "partially ruled multi-panel table", "filled AcroForm with comments",
                                           "rotated pages"],
            },
            "xlsx": {
                "available": True,
                "options": {"include_images": [True, False], "include_hyperlinks": [True, False],
                            "include_form_values": [True, False], "include_comments": [True, False],
                            "running_headers": ["preserve"], "ocr": ["off"], "locale": [None],
                            "allow_partial": [True, False]},
                "unsupported_options": {"layout_mode": "not applicable to XLSX", "ocr": "OCR is not implemented",
                                        "locale": "no locale rules implemented"},
                "verified_reader": "Microsoft Excel for Mac",
                "verified_fixture_scope": ["ruled multi-page lookup table (IRS EIC)",
                                           "ruled header over unruled body with row labels (EIA WPSR table 1)",
                                           "ruled header with row labels outside the grid and footnote marks (Census P60 A-4b)"],
                "fidelity_limits": XLSX_FIDELITY_LIMITS,
            },
            "html": {
                "available": True,
                "options": {"include_images": [True, False], "include_hyperlinks": [True, False],
                            "include_form_values": [True, False], "include_comments": [True, False],
                            "running_headers": ["preserve"], "layout_mode": ["preserve", "reflow"],
                            "ocr": ["off"], "locale": [None], "allow_partial": [True, False]},
                "modes": {"preserve": {"name": "Preserve layout", "default": True,
                                       "fidelity_limits": HTML_FIDELITY_LIMITS["preserve"]},
                          "reflow": {"name": "Responsive reading", "default": False,
                                     "fidelity_limits": HTML_FIDELITY_LIMITS["reflow"]}},
                "unsupported_options": {"ocr": "OCR is not implemented", "locale": "no locale rules implemented"},
                "verified_reader": "Google Chrome for Mac",
            },
            "md": {
                "available": True,
                "dialect": "gfm",
                "options": {"include_images": [True, False], "include_hyperlinks": [True, False],
                            "include_form_values": [True, False], "include_comments": [True, False],
                            "running_headers": ["preserve"], "ocr": ["off"], "locale": [None],
                            "allow_partial": [True, False], "markdown_images": ["folder", "embed"],
                            "markdown_asset_folder": ["<plain folder name>"]},
                "markdown_images": {"folder": "default: pictures are files in a sibling folder named by markdown_asset_folder (default 'document_images'), linked relatively; listed in artifact.files",
                                    "embed": "pictures are data URIs inside the one file (some renderers, GitHub among them, do not show them)"},
                "unsupported_options": {"layout_mode": "not applicable to Markdown", "ocr": "OCR is not implemented",
                                        "locale": "no locale rules implemented"},
                "verified_reader": "pandoc 3 GFM parser, rendered in Google Chrome",
                "fidelity_limits": MD_FIDELITY_LIMITS,
            },
            "pptx": {
                "available": True,
                "options": {"include_images": [True, False], "include_hyperlinks": [True, False],
                            "include_form_values": [True, False], "include_comments": [True, False],
                            "running_headers": ["preserve"], "ocr": ["off"], "locale": [None],
                            "allow_partial": [True, False], "pptx_mode": ["editable", "page_image"]},
                "modes": {"editable": {"name": "Editable slides", "default": True,
                                       "fidelity_limits": PPTX_FIDELITY_LIMITS["editable"]},
                          "page_image": {"name": "Page image with editable text", "default": False,
                                         "fidelity_limits": PPTX_FIDELITY_LIMITS["page_image"]}},
                "unsupported_options": {"layout_mode": "not applicable to PPTX (slides are always laid out as the pages)",
                                        "ocr": "OCR is not implemented", "locale": "no locale rules implemented"},
                "verified_reader": "Microsoft PowerPoint for Mac",
            },
            "rtf": {
                "available": True,
                "options": {"include_images": [True, False], "include_hyperlinks": [True, False],
                            "include_form_values": [True, False], "include_comments": [True, False],
                            "running_headers": ["preserve"], "ocr": ["off"], "locale": [None],
                            "allow_partial": [True, False]},
                "unsupported_options": {"layout_mode": "not applicable to RTF (reading order only)",
                                        "ocr": "OCR is not implemented", "locale": "no locale rules implemented"},
                "verified_reader": "Microsoft Word for Mac and TextEdit",
                "fidelity_limits": RTF_FIDELITY_LIMITS,
            },
            "xml": {
                "available": True,
                "schema": {"namespace": "urn:zpdf:export:document:1", "version": "1",
                           "file": "zpdf-document-1.xsd", "documentation": "docs/xml-format.md"},
                "options": {"include_images": [True, False], "include_hyperlinks": [True, False],
                            "include_form_values": [True, False], "include_comments": [True, False],
                            "running_headers": ["preserve"], "ocr": ["off"], "locale": [None],
                            "allow_partial": [True, False], "xml_images": ["embed", "omit"]},
                "unsupported_options": {"layout_mode": "not applicable to XML",
                                        "ocr": "OCR is not implemented", "locale": "no locale rules implemented"},
                "verified_reader": "the shipped XSD (lxml) and xmllint",
                "fidelity_limits": XML_FIDELITY_LIMITS,
            },
            "epub": {
                "available": True,
                "version": "3",
                "options": {"include_images": [True, False], "include_hyperlinks": [True, False],
                            "include_form_values": [True, False], "include_comments": [True, False],
                            "running_headers": ["preserve"], "ocr": ["off"], "locale": [None],
                            "allow_partial": [True, False]},
                "unsupported_options": {"layout_mode": "not applicable to EPUB (reflowable only)",
                                        "ocr": "OCR is not implemented", "locale": "no locale rules implemented"},
                "verified_reader": "epubcheck 4.2.6 (verification only)",
                "fidelity_limits": EPUB_FIDELITY_LIMITS,
            },
        },
        "ocr": {"available": False, "reason": "not_implemented"},
        "languages": {"qualified": ["und"], "note": "no language-specific behavior; left-to-right horizontal text only"},
        "directions": ["ltr"],
        "resource_limits": {"enforced_by_worker": [], "note": "limits are the coordinator's responsibility in this increment"},
    }


class _Cancelled(Exception):
    pass


class Worker:
    def __init__(self, stdin=None, stdout=None, stderr=None) -> None:
        self.stdin = stdin or sys.stdin
        self.stdout = stdout or sys.stdout
        self.stderr = stderr or sys.stderr
        self.lines: queue.Queue[str | None] = queue.Queue()
        self.deferred: list[str] = []
        self.cancel_job: str | None = None

    # -- transport -------------------------------------------------------
    def emit(self, obj: dict) -> None:
        self.stdout.write(json.dumps(obj, separators=(",", ":"), ensure_ascii=False) + "\n")
        self.stdout.flush()

    def log(self, msg: str) -> None:
        self.stderr.write(msg[:200] + "\n")
        self.stderr.flush()

    def _reader(self) -> None:
        for line in self.stdin:
            self.lines.put(line)
        self.lines.put(None)

    def run(self) -> None:
        threading.Thread(target=self._reader, daemon=True).start()
        while True:
            if self.deferred:
                line = self.deferred.pop(0)
            else:
                line = self.lines.get()
            if line is None:
                return
            if not line.strip():
                continue
            self.handle(line)

    # -- dispatch --------------------------------------------------------
    def handle(self, line: str) -> None:
        try:
            req = json.loads(line)
        except json.JSONDecodeError:
            self._error_result(None, ExportError("INVALID_REQUEST", "malformed_json"))
            return
        if not isinstance(req, dict):
            self._error_result(None, ExportError("INVALID_REQUEST", "request_not_object"))
            return
        job_id = req.get("job_id") if isinstance(req.get("job_id"), str) else None
        try:
            if req.get("protocol_version") != PROTOCOL_VERSION:
                raise ExportError("PROTOCOL_MISMATCH", "unsupported_protocol_version",
                                  supported=PROTOCOL_VERSION)
            op = req.get("operation")
            if op == "capabilities":
                self.emit(capabilities())
            elif op == "convert":
                self.convert(req)
            elif op == "cancel":
                # a cancel arriving when no job is running is a no-op acknowledgement
                self.emit({"protocol_version": PROTOCOL_VERSION, "event": "cancel_ack",
                           "job_id": job_id, "active": False})
            else:
                raise ExportError("INVALID_REQUEST", "unknown_operation")
        except ExportError as exc:
            self._error_result(job_id, exc)
        except Exception as exc:  # noqa: BLE001
            self.log(f"WORKER_FAILED {type(exc).__name__}")
            self._error_result(job_id, ExportError("WORKER_FAILED", "unexpected_exception",
                                                   exception=type(exc).__name__))

    def _error_result(self, job_id: str | None, exc: ExportError) -> None:
        status = "cancelled" if exc.code == "CANCELLED" else "error"
        out = {"protocol_version": PROTOCOL_VERSION, "event": "result", "job_id": job_id,
               "status": status, "backend": BACKEND_ID}
        if status == "error":
            out["error"] = exc.to_dict()
        self.emit(out)

    # -- cancellation ----------------------------------------------------
    def _poll_cancel(self, job_id: str) -> bool:
        while True:
            try:
                line = self.lines.get_nowait()
            except queue.Empty:
                break
            if line is None:
                self.lines.put(None)
                return True  # stdin closed: coordinator is gone, stop work
            try:
                msg = json.loads(line)
            except json.JSONDecodeError:
                self.deferred.append(line)
                continue
            if isinstance(msg, dict) and msg.get("operation") == "cancel" and msg.get("job_id") == job_id:
                self.cancel_job = job_id
            else:
                self.deferred.append(line)
        return self.cancel_job == job_id

    # -- convert ---------------------------------------------------------
    def convert(self, req: dict) -> None:
        job_id = req.get("job_id")
        if not isinstance(job_id, str) or not job_id:
            raise ExportError("INVALID_REQUEST", "job_id_missing")
        unknown = set(req) - _KNOWN_REQUEST_KEYS
        if unknown:
            raise ExportError("INVALID_REQUEST", "unknown_request_field", fields=sorted(unknown))
        snap = req.get("snapshot")
        if not isinstance(snap, dict) or not isinstance(snap.get("path"), str) or not isinstance(snap.get("sha256"), str):
            raise ExportError("INVALID_REQUEST", "snapshot_missing_fields")
        staging = req.get("staging_directory")
        if not isinstance(staging, str):
            raise ExportError("INVALID_REQUEST", "staging_directory_missing")
        staging_path = Path(staging)
        if not staging_path.is_dir() or staging_path.is_symlink():
            raise ExportError("INVALID_REQUEST", "staging_directory_invalid")
        if any(staging_path.iterdir()):
            raise ExportError("INVALID_REQUEST", "staging_directory_not_empty")
        fmt = req.get("format")
        if fmt not in _ARTIFACT_NAMES:
            reason = "unknown_format"
            raise ExportError("UNSUPPORTED_FORMAT", reason, format=str(fmt))
        opts_in = req.get("options", {})
        if not isinstance(opts_in, dict):
            raise ExportError("INVALID_REQUEST", "options_not_object")
        unknown_opts = set(opts_in) - _KNOWN_OPTIONS
        if unknown_opts:
            raise ExportError("UNSUPPORTED_OPTION", "unknown_option", option=sorted(unknown_opts)[0])
        not_applicable = set(opts_in) - _FORMAT_OPTIONS[fmt]
        if not_applicable:
            raise ExportError("UNSUPPORTED_OPTION", "option_not_applicable_to_format",
                              option=sorted(not_applicable)[0], format=fmt)
        opts = dict(opts_in)
        md_images = opts.pop("markdown_images", "folder")
        md_folder = opts.pop("markdown_asset_folder", "document_images")
        pptx_mode = opts.pop("pptx_mode", "editable")
        xml_images = opts.pop("xml_images", "embed")
        if fmt == "xml" and xml_images not in ("embed", "omit"):
            raise ExportError("UNSUPPORTED_OPTION", "xml_images_unknown", option="xml_images", value=str(xml_images))
        if fmt == "pptx" and pptx_mode not in ("editable", "page_image"):
            raise ExportError("UNSUPPORTED_OPTION", "pptx_mode_unknown", option="pptx_mode", value=str(pptx_mode))
        if fmt == "md":
            if md_images not in ("folder", "embed"):
                raise ExportError("UNSUPPORTED_OPTION", "markdown_images_unknown", option="markdown_images", value=str(md_images))
            if not isinstance(md_folder, str) or not _ASSET_FOLDER.match(md_folder) or md_folder.rstrip(" .") != md_folder:
                raise ExportError("INVALID_REQUEST", "markdown_asset_folder_invalid")
        if fmt == "xlsx":
            opts.update(layout_mode="reflow", data_tables=True)   # no page layouts; recover unruled table bodies
        elif fmt == "md":
            opts.update(layout_mode="reflow", data_tables=True)
        elif fmt == "rtf":
            opts.update(layout_mode="reflow", data_tables=True)
        elif fmt == "xml":
            opts.update(layout_mode="reflow", data_tables=True, keep_geometry=True)
        elif fmt == "epub":
            opts.update(layout_mode="reflow", data_tables=True)
        elif fmt == "pptx":
            opts.update(layout_mode="preserve", page_drawings=(pptx_mode == "page_image"))
        elif fmt == "html" and opts.get("layout_mode") == "reflow":
            opts.update(data_tables=True)                         # responsive tables carry recovered bodies
        options = ExportOptions(**opts)
        pages = req.get("pages")

        self.cancel_job = None
        artifact_name = _ARTIFACT_NAMES[fmt]
        artifact = staging_path / artifact_name
        part = staging_path / (artifact_name + ".part")

        def progress(stage: str, done: int, total: int) -> None:
            self.emit({"protocol_version": PROTOCOL_VERSION, "event": "progress", "job_id": job_id,
                       "stage": stage, "pages_done": done, "pages_total": total})

        def cancelled() -> bool:
            return self._poll_cancel(job_id)

        try:
            progress("extract", 0, len(pages) if isinstance(pages, list) else 0)
            doc = build_document(Path(snap["path"]), snap["sha256"], pages, options, progress, cancelled)
            if cancelled():
                raise ExportError("CANCELLED", "cancelled_before_write")
            progress("write", 0, 1)
            if fmt == "xlsx":
                from .writers.xlsx_writer import write_xlsx
                wstats = write_xlsx(doc, part)
            elif fmt == "html":
                from .writers.html_writer import write_html
                wstats = write_html(doc, part, options.layout_mode)
            elif fmt == "md":
                from .writers.markdown_writer import write_markdown
                wstats = write_markdown(doc, part, images=md_images, asset_folder=md_folder)
            elif fmt == "pptx":
                from .writers.pptx_writer import write_pptx
                wstats = write_pptx(doc, part, mode=pptx_mode)
            elif fmt == "rtf":
                from .writers.rtf_writer import write_rtf
                wstats = write_rtf(doc, part)
            elif fmt == "xml":
                from .writers.xml_writer import write_xml
                wstats = write_xml(doc, part, images=xml_images)
            elif fmt == "epub":
                from .writers.epub_writer import write_epub
                wstats = write_epub(doc, part)
            elif options.layout_mode == "preserve":
                from .writers.docx_layout_writer import write_docx_layout
                wstats = write_docx_layout(doc, part)
            else:
                wstats = write_docx(doc, part)
            if cancelled():
                raise ExportError("CANCELLED", "cancelled_before_validate")
            progress("write", 1, 1)
            progress("validate", 0, 1)
            if fmt == "md":
                _validate_md(part, staging_path, wstats.get("files", []))
            else:
                {"xlsx": _validate_xlsx, "html": _validate_html, "pptx": _validate_pptx,
                 "rtf": _validate_rtf, "xml": _validate_xml,
                 "epub": _validate_epub}.get(fmt, _validate_docx)(part)
            os.replace(part, artifact)
            progress("validate", 1, 1)
            digest = hashlib.sha256(artifact.read_bytes()).hexdigest()
            warnings = list(doc.warnings) + list(wstats.pop("warnings", []))
            if wstats.get("images_downsampled"):
                # embedding below the source resolution lowers quality: disclose it per page
                sel = list(pages) if isinstance(pages, list) else []
                by_section = wstats.get("images_downsampled_by_section") or {}
                for sec in wstats.get("images_downsampled_pages") or []:
                    warnings.append({"code": "IMAGE_DOWNSAMPLED", "page": sel[sec - 1] if 0 < sec <= len(sel) else sec,
                                     "count": by_section.get(sec, 0),
                                     "max_dpi": wstats.get("max_picture_dpi"),
                                     "detail": "pictures embedded below their source resolution for the size they are placed at"})
            if wstats.get("links_dropped_unsafe_scheme"):
                warnings.append({"code": "LINK_SCHEME_NOT_ALLOWED", "count": wstats["links_dropped_unsafe_scheme"]})
            self.emit({
                "protocol_version": PROTOCOL_VERSION, "event": "result", "job_id": job_id,
                "status": "ok_with_warnings" if warnings else "ok",
                "backend": BACKEND_ID, "snapshot_sha256": doc.source_sha256,
                "layout_mode": options.layout_mode if fmt in ("docx", "html") else None,
                **({"pptx_mode": pptx_mode} if fmt == "pptx" else {}),
                "artifact": {"relative_path": artifact_name, "sha256": digest, "bytes": artifact.stat().st_size,
                             **({"files": wstats.pop("files")} if wstats.get("files") else {})},
                "warnings": warnings,
                "stats": {**doc.stats, "writer": {k: v for k, v in wstats.items() if k != "font_substitution"}},
                "report": {"versions": versions(), "pages": list(pages), "format": fmt,
                           "options": opts_in, "font_substitution": wstats.get("font_substitution", "mapped from PDF font names"),
                           "font_mapping": dict(sorted((wstats.get("fonts") or {}).items()))},
            })
        except BaseException:
            for p in (part, artifact):
                try:
                    p.unlink()
                except FileNotFoundError:
                    pass
            if fmt == "md":                      # the picture folder goes with the document
                import shutil
                shutil.rmtree(staging_path / md_folder, ignore_errors=True)
            raise


def _validate_md(path: Path, staging: Path | None = None, files: list | None = None) -> None:
    """UTF-8 Markdown with no script element (raw HTML in Markdown is passed
    through by many renderers); every relative picture link points at a listed
    file in staging, and every listed file is linked and intact."""
    import hashlib
    import re
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise ExportError("OUTPUT_INVALID", "md_not_utf8") from exc
    if re.search(r"(?i)<script", text):
        raise ExportError("OUTPUT_INVALID", "md_contains_script")
    if staging is None:
        return
    listed = {f["relative_path"]: f for f in files or []}
    links = {m for m in re.findall(r"!\[[^\]]*\]\(<?([^)>]+?)>?\)", text) if not m.startswith("data:")}
    if links != set(listed):
        raise ExportError("OUTPUT_INVALID", "md_picture_links_mismatch")
    for rel, f in listed.items():
        p = staging / rel
        if p.is_symlink() or not p.is_file() or ".." in Path(rel).parts:
            raise ExportError("OUTPUT_INVALID", "md_picture_missing")
        if hashlib.sha256(p.read_bytes()).hexdigest() != f["sha256"]:
            raise ExportError("OUTPUT_INVALID", "md_picture_hash_mismatch")


def _validate_html(path: Path) -> None:
    """A self-contained HTML document: UTF-8, a doctype, no script, and no
    resource loaded from outside the file (links to web pages are allowed)."""
    import re
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise ExportError("OUTPUT_INVALID", "html_not_utf8") from exc
    if not text.startswith("<!doctype html>") or "</html>" not in text:
        raise ExportError("OUTPUT_INVALID", "html_incomplete")
    if re.search(r"(?i)<script", text):
        raise ExportError("OUTPUT_INVALID", "html_contains_script")
    if re.search(r'(?i)\b(src|srcset)\s*=\s*"(?!data:)', text) or re.search(r"(?i)url\(\s*['\"]?(?!data:)", text):
        raise ExportError("OUTPUT_INVALID", "html_external_resource")


def _validate_xlsx(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as z:
            if z.testzip() is not None:
                raise ExportError("OUTPUT_INVALID", "zip_member_corrupt")
            names = set(z.namelist())
        for required in ("[Content_Types].xml", "xl/workbook.xml", "_rels/.rels"):
            if required not in names:
                raise ExportError("OUTPUT_INVALID", "xlsx_part_missing", part=required)
        from openpyxl import load_workbook
        with open(path, "rb") as fh:          # a file object: the staged name ends in .part
            wb = load_workbook(fh)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.data_type == "f":
                        raise ExportError("OUTPUT_INVALID", "xlsx_contains_formula", sheet=ws.title)
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExportError("OUTPUT_INVALID", "xlsx_reopen_failed", exception=type(exc).__name__) from exc


def _validate_epub(path: Path) -> None:
    """The EPUB container rules the book must meet: mimetype first, stored, with
    no extra field; container.xml naming a package document; every manifest
    item present and every file declared; every content document well-formed
    XHTML without script; no resource loaded from outside the package."""
    import re
    from lxml import etree
    try:
        with zipfile.ZipFile(path) as z:
            if z.testzip() is not None:
                raise ExportError("OUTPUT_INVALID", "zip_member_corrupt")
            infos = z.infolist()
            if not infos or infos[0].filename != "mimetype" or infos[0].compress_type != zipfile.ZIP_STORED \
                    or infos[0].extra or z.read("mimetype") != b"application/epub+zip":
                raise ExportError("OUTPUT_INVALID", "epub_mimetype_invalid")
            parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
            container = etree.fromstring(z.read("META-INF/container.xml"), parser)
            root = container.find(".//{urn:oasis:names:tc:opendocument:xmlns:container}rootfile")
            opf_path = root.get("full-path")
            opf = etree.fromstring(z.read(opf_path), parser)
            base = opf_path.rsplit("/", 1)[0] + "/" if "/" in opf_path else ""
            ns = {"o": "http://www.idpf.org/2007/opf"}
            items = opf.findall(".//o:manifest/o:item", ns)
            declared = {base + i.get("href") for i in items}
            names = set(z.namelist())
            if not declared <= names or names - declared - {"mimetype", "META-INF/container.xml", opf_path}:
                raise ExportError("OUTPUT_INVALID", "epub_manifest_mismatch")
            for i in items:
                if i.get("media-type") == "application/xhtml+xml":
                    data = z.read(base + i.get("href"))
                    etree.fromstring(data, parser)
                    text = data.decode("utf-8")
                    if re.search(r"(?i)<script", text):
                        raise ExportError("OUTPUT_INVALID", "epub_contains_script")
                    if re.search(r'(?i)\b(src|srcset)\s*=\s*"[a-z]+:', text) or re.search(r"(?i)url\(\s*['\"]?[a-z]+:", text):
                        raise ExportError("OUTPUT_INVALID", "epub_external_resource")
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExportError("OUTPUT_INVALID", "epub_reopen_failed", exception=type(exc).__name__) from exc


def _validate_xml(path: Path) -> None:
    """Well-formed UTF-8 XML, no DTD or entity declarations, valid against the
    shipped schema (zpdf-document-1.xsd)."""
    from lxml import etree
    from .writers.xml_writer import schema
    try:
        parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
        tree = etree.parse(str(path), parser)
    except etree.XMLSyntaxError as exc:
        raise ExportError("OUTPUT_INVALID", "xml_not_well_formed") from exc
    if tree.docinfo.internalDTD is not None or tree.docinfo.doctype:
        raise ExportError("OUTPUT_INVALID", "xml_has_doctype")
    xs = schema()
    if not xs.validate(tree):
        raise ExportError("OUTPUT_INVALID", "xml_schema_invalid")   # the error text may quote document text


def _validate_rtf(path: Path) -> None:
    """7-bit RTF that opens with the RTF 1 header, has balanced groups, and holds
    no embedded object and no field but HYPERLINK to http, https or mailto (an
    INCLUDEPICTURE or INCLUDETEXT field would fetch content when opened)."""
    import re
    data = path.read_bytes()
    if any(b > 127 for b in data):
        raise ExportError("OUTPUT_INVALID", "rtf_not_7bit")
    text = data.decode("ascii")
    if not text.startswith("{\\rtf1"):
        raise ExportError("OUTPUT_INVALID", "rtf_header_missing")
    depth, i = 0, 0
    while i < len(text):
        c = text[i]
        if c == "\\":
            i += 2
            continue
        depth += (c == "{") - (c == "}")
        if depth < 0:
            break
        i += 1
    if depth != 0:
        raise ExportError("OUTPUT_INVALID", "rtf_groups_unbalanced")
    if re.search(r"\\(object|objdata|objemb|objlink|objautlink)\b", text):
        raise ExportError("OUTPUT_INVALID", "rtf_embedded_object")
    for inst in re.findall(r"\\fldinst\s*\{?\s*([^}]*)", text):
        if not re.match(r'HYPERLINK "(https?:|mailto:)', inst.strip()):
            raise ExportError("OUTPUT_INVALID", "rtf_field_not_allowed")


def _validate_pptx(path: Path) -> None:
    """A presentation python-pptx reopens, with no macro, ActiveX or other
    binary part, and nothing outside the package but hyperlink targets."""
    import re
    try:
        with zipfile.ZipFile(path) as z:
            if z.testzip() is not None:
                raise ExportError("OUTPUT_INVALID", "zip_member_corrupt")
            names = set(z.namelist())
            for required in ("[Content_Types].xml", "ppt/presentation.xml", "_rels/.rels"):
                if required not in names:
                    raise ExportError("OUTPUT_INVALID", "pptx_part_missing", part=required)
            if any(n.endswith(".bin") or "vbaProject" in n or "activeX" in n for n in names):
                raise ExportError("OUTPUT_INVALID", "pptx_binary_part")
            for n in names:
                if n.endswith(".rels"):
                    for rel in re.findall(r"<Relationship\b[^>]*>", z.read(n).decode("utf-8", "replace")):
                        if 'TargetMode="External"' in rel and not rel.count("/hyperlink\""):
                            raise ExportError("OUTPUT_INVALID", "pptx_external_resource")
        from pptx import Presentation
        with open(path, "rb") as fh:          # a file object: the staged name ends in .part
            Presentation(fh)
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExportError("OUTPUT_INVALID", "pptx_reopen_failed", exception=type(exc).__name__) from exc


def _validate_docx(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as z:
            if z.testzip() is not None:
                raise ExportError("OUTPUT_INVALID", "zip_member_corrupt")
            names = set(z.namelist())
        for required in ("[Content_Types].xml", "word/document.xml", "_rels/.rels"):
            if required not in names:
                raise ExportError("OUTPUT_INVALID", "docx_part_missing", part=required)
        from docx import Document as DocxDocument
        DocxDocument(str(path))
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExportError("OUTPUT_INVALID", "docx_reopen_failed", exception=type(exc).__name__) from exc


def main() -> None:
    Worker().run()


if __name__ == "__main__":
    main()
