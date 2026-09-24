"""DOCX writer over the ordered IR using python-docx (MIT) + lxml (BSD).

Editable paragraphs, heading/list styles, true tables with rectangular spans,
inline images, hyperlinks (http/https/mailto only), and labeled form-value and
comment sections anchored by source page.
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from urllib.parse import urlsplit

from docx import Document as DocxDocument
from docx.enum.text import WD_COLOR_INDEX  # noqa: F401  (kept for clarity of imports)
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..ir import CommentNode, Document, FigureNode, FormValueNode, TableNode, TextNode

_ALLOWED_SCHEMES = {"http", "https", "mailto"}
_MAX_IMAGE_WIDTH_PT = 468  # 6.5in text width on US Letter with 1in margins


def write_docx(doc: Document, path: Path) -> dict:
    d = DocxDocument()
    cp = d.core_properties
    cp.author = "zPDF Export"
    cp.last_modified_by = "zPDF Export"
    cp.title = ""
    cp.comments = ""
    stats = {"paragraphs": 0, "tables": 0, "images": 0, "form_values": 0, "comments": 0,
             "links_dropped_unsafe_scheme": 0, "font_substitution": "template defaults (Calibri/Cambria)"}
    section_page = {"form_value": None, "comment": None}
    pending_values: list[FormValueNode] = []

    def flush_values():
        if pending_values:
            _write_form_values(d, pending_values, section_page["form_value"])
            stats["form_values"] += len(pending_values)
            pending_values.clear()

    for nid in doc.flow:
        node = doc.nodes[nid]
        page = node.source_regions[0]["page"] if node.source_regions else None
        if isinstance(node, FormValueNode):
            if node.placement in ("table_cell", "inline_sentence"):
                stats["form_values"] += 1  # value already written inside its table cell / sentence
                continue
            if node.placement == "inline" and node.paired:
                flush_values()
                _write_inline_field(d, node)
                stats["form_values"] += 1
                continue
            if section_page["form_value"] != page:
                flush_values()
                section_page["form_value"] = page
            pending_values.append(node)
            continue
        flush_values()
        if isinstance(node, TextNode):
            _write_text(d, node, stats)
            stats["paragraphs"] += 1
            continue
        stats["_list_open"] = None
        if isinstance(node, TableNode):
            _write_table(d, node)
            stats["tables"] += 1
        elif isinstance(node, FigureNode):
            asset = doc.assets[node.asset_id]
            width_pt = min(node.source_regions[0]["bbox"][2] - node.source_regions[0]["bbox"][0],
                           _MAX_IMAGE_WIDTH_PT) if node.source_regions else _MAX_IMAGE_WIDTH_PT
            width_pt = max(width_pt, 8)
            pic = d.add_picture(io.BytesIO(asset.data), width=Pt(width_pt))
            if node.absorbed_text:
                pic._inline.docPr.set("descr", node.absorbed_text)   # the labels the picture keeps
            stats["images"] += 1
        elif isinstance(node, CommentNode):
            if section_page["comment"] != page:
                section_page["comment"] = page
                d.add_heading(f"Comments (page {page})", level=3)
            p = d.add_paragraph()
            lead = p.add_run(f"{node.comment_kind.capitalize()}")
            lead.bold = True
            if node.author:
                p.add_run(f" by {node.author}")
            if node.modified:
                p.add_run(f" ({node.modified})")
            p.add_run(": ")
            p.add_run(node.text)
            stats["comments"] += 1
    flush_values()
    stats.pop("_list_open", None)
    d.save(str(path))
    return stats


_ENUM_START = re.compile(r"^\(?(1|a|A|i|I)[.)]$")


def _new_numbering_instance(d) -> int | None:
    """Create a fresh w:num for the List Number style so numbering restarts at 1."""
    try:
        style = d.styles["List Number"]
        num_id = style.element.pPr.numPr.numId.val
        numbering = d.part.numbering_part.element
        base = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(num_id))
        abstract_id = base.find(qn("w:abstractNumId")).get(qn("w:val"))
        existing = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        new_id = max(existing) + 1
        num = OxmlElement("w:num"); num.set(qn("w:numId"), str(new_id))
        an = OxmlElement("w:abstractNumId"); an.set(qn("w:val"), abstract_id); num.append(an)
        ov = OxmlElement("w:lvlOverride"); ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); ov.append(so); num.append(ov)
        numbering.append(num)
        return new_id
    except Exception:  # noqa: BLE001 - fall back to the style's shared numbering
        return None


def _apply_numbering(p, num_id: int) -> None:
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); numPr.append(ilvl)
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id)); numPr.append(nid)
    pPr.append(numPr)


def _write_text(d, node: TextNode, stats: dict) -> None:
    if node.kind == "invisible":
        # the page's invisible text layer: searchable hidden text, as the source has it
        p = d.add_paragraph()
        r = p.add_run(node.text.replace("\n", " "))
        r.font.hidden = True
        return
    numbered = node.kind == "list_item" and bool(node.marker and node.marker[0].isalnum())
    if not numbered:
        stats["_list_open"] = None  # any non-numbered content ends the open list
    if node.kind == "heading":
        p = d.add_heading("", level=min(max(node.level or 1, 1), 9))
    elif node.kind == "list_item":
        p = d.add_paragraph(style="List Number" if numbered else "List Bullet")
        if numbered:
            # a marker that starts a sequence, or the first numbered item after
            # non-list content, begins a new list with its own numbering instance
            if stats.get("_list_open") is None or _ENUM_START.match(node.marker or ""):
                stats["_list_open"] = _new_numbering_instance(d)
            if stats["_list_open"] is not None:
                _apply_numbering(p, stats["_list_open"])
    elif node.kind == "unsupported":
        p = d.add_paragraph()
        r = p.add_run(node.text)
        r.italic = True
        return
    else:
        p = d.add_paragraph()
    runs = node.runs or [{"text": node.text, "bold": False, "italic": False}]
    for r in runs:
        uri = r.get("uri")
        if uri and urlsplit(uri).scheme.lower() in _ALLOWED_SCHEMES:
            _add_hyperlink(p, uri, r["text"], r.get("bold", False), r.get("italic", False))
        else:
            if uri:
                stats["links_dropped_unsafe_scheme"] += 1
            run = p.add_run(r["text"])
            run.bold = bool(r.get("bold")) or None
            run.italic = bool(r.get("italic")) or None
            if node.kind == "heading" and r.get("bold"):
                run.bold = True


def _add_hyperlink(paragraph, url: str, text: str, bold: bool, italic: bool) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rPr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    if bold:
        rPr.append(OxmlElement("w:b"))
    if italic:
        rPr.append(OxmlElement("w:i"))
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _write_table(d, node: TableNode) -> None:
    table = d.add_table(rows=node.n_rows, cols=node.n_cols)
    table.style = "Table Grid"
    header = set(node.header_rows)
    for c in node.cells:
        r, col = c["row"], c["col"]
        cell = table.cell(r, col)
        rs, cs = c.get("rowspan", 1), c.get("colspan", 1)
        if rs > 1 or cs > 1:
            cell = cell.merge(table.cell(r + rs - 1, col + cs - 1))
        text = c["raw_text"]
        para = cell.paragraphs[0]
        run = para.add_run(text)
        if c.get("bold") or r in header:
            run.bold = True
    d.add_paragraph()


def _state_glyph(v: FormValueNode) -> str:
    return "\u2612" if v.checked else "\u2610"


def _write_inline_field(d, v: FormValueNode) -> None:
    p = d.add_paragraph()
    if v.field_type in ("checkbox", "radio"):
        p.add_run(_state_glyph(v) + " ")
        label = v.label or v.field_name
        p.add_run(label)
        ev = (v.export_value or "").strip()
        if v.checked and ev and ev.lower() not in ("on", "yes", "true", "1") and ev.lower() not in label.lower():
            p.add_run(f" ({ev})")
        return
    lead = p.add_run((v.label or v.field_name).rstrip(":") + ": ")
    lead.bold = True
    if v.raw_value.strip():
        p.add_run(v.raw_value)
    else:
        r = p.add_run("(blank)")
        r.italic = True


def _write_form_values(d, values: list[FormValueNode], page: int | None) -> None:
    """Fallback section: values whose visible label could not be determined."""
    d.add_heading(f"Unpaired form values (page {page})", level=3)
    note = d.add_paragraph()
    r = note.add_run("These field values could not be matched to a visible label with confidence. "
                     "Field names and tooltips come from the PDF form definition.")
    r.italic = True
    table = d.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(("Field name", "Tooltip", "Value", "Type")):
        run = table.cell(0, i).paragraphs[0].add_run(h)
        run.bold = True
    for v in values:
        row = table.add_row().cells
        row[0].paragraphs[0].add_run(v.field_name or "(unnamed field)")
        row[1].paragraphs[0].add_run(v.tooltip or "")
        value = v.raw_value
        if v.checked is not None:
            value = f"{_state_glyph(v)} {v.export_value or ''}".strip() + (" (checked)" if v.checked else " (unchecked)")
        row[2].paragraphs[0].add_run(value)
        kind = v.field_type
        if v.ambiguity:
            kind += " — candidates: " + " | ".join(v.ambiguity[:3])
        row[3].paragraphs[0].add_run(kind)
    d.add_paragraph()
