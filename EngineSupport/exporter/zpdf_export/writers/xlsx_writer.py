"""XLSX writer: one worksheet per table (tables that continue across page
halves or pages are joined), merged cells for spans, header rows bold and
frozen, and a "Text" sheet carrying every non-table element in reading order.

Cell typing never changes what a cell shows: a cell becomes a number only when
the number, written with the chosen number format, reads exactly as printed
(grouping, decimals and percent included). Leading-zero identifiers, dates,
parenthesised labels, footnote-marked values and anything ambiguous stay
text. No formula is ever written; text that begins with "=" is a string.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..ir import CommentNode, FigureNode, FormValueNode, TableNode, TextNode

_NUM = re.compile(r"^(-)?(\d{1,3}(?:,\d{3})+|\d+)(?:\.(\d+))?(%)?$")
_MINUS = "−"


def display(value, fmt: str) -> str:
    """The text Excel shows for a number in one of the formats cell_value
    chooses (General, 0, 0.0…, #,##0…, …%)."""
    if fmt in ("General", "0"):
        return str(value)
    pct = fmt.endswith("%")
    body = fmt[:-1] if pct else fmt
    dec = len(body.split(".")[1]) if "." in body else 0
    grouped = body.startswith("#,##0")
    v = value * 100 if pct else value
    text = f"{v:,.{dec}f}" if grouped else f"{v:.{dec}f}"
    return text + ("%" if pct else "")


def cell_value(raw: str):
    """(value, number_format, kind) for a cell's raw text. kind is "number"
    or "text"; a text value is the raw text unchanged."""
    s = raw.strip() if isinstance(raw, str) else ""
    if not s or "\n" in s:
        return raw, "General", "text"
    t = "-" + s[1:] if s.startswith(_MINUS) else s
    m = _NUM.match(t)
    if not m:
        return raw, "General", "text"
    neg, whole, frac, pct = m.groups()
    digits = whole.replace(",", "")
    if len(digits) > 1 and digits.startswith("0"):
        return raw, "General", "text"          # an identifier (ZIP code, account, line number)
    if len(digits) + len(frac or "") > 15:
        return raw, "General", "text"          # beyond Excel's precision: kept exactly as text
    grouped = "," in whole
    dec = len(frac or "")
    if dec:
        number = float(("-" if neg else "") + digits + "." + frac)
    else:
        number = int(("-" if neg else "") + digits)
    body = ("#,##0" if grouped else "0") + ("." + "0" * dec if dec else "")
    if pct:
        value, fmt = round(number / 100, dec + 6), body + "%"
    elif not grouped and not dec:
        value, fmt = number, ("General" if len(digits) <= 11 else "0")
    else:
        value, fmt = number, body
    if display(value, fmt) != t:
        return raw, "General", "text"          # the number would not read as printed: keep the text
    return value, fmt, "number"


def _clean(text: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("�", text)


def _put(cell, raw: str) -> str:
    """Write raw text into a cell as a number (when it displays identically)
    or as a string. Returns the kind written."""
    value, fmt, kind = cell_value(raw)
    if kind == "number":
        cell.value = value
        cell.number_format = fmt
    else:
        cell.value = _clean(raw)
        cell.data_type = "s"                    # never a formula, whatever it starts with
    return kind


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("–", "-").replace("—", "-")).strip().lower()


def _header_key(t: TableNode):
    return tuple(sorted((c["row"], c["col"], c["rowspan"], c["colspan"], _norm(c["raw_text"]))
                        for c in t.cells if c["row"] in t.header_rows))


def _joinable(a: TableNode, b: TableNode) -> bool:
    """A table continues another (the next half of a page, or the next page)
    when both have the same columns and the same non-empty header rows."""
    return (a.n_cols == b.n_cols and a.header_rows and a.header_rows == b.header_rows
            and a.header_rows == list(range(len(a.header_rows))) and _header_key(a) == _header_key(b))


def _row_key(t: TableNode, row: int):
    return tuple(sorted((c["col"], c["colspan"], _norm(c["raw_text"])) for c in t.cells if c["row"] == row))


def _repeated_rows(first: TableNode, cont: TableNode) -> int:
    """Leading rows of a continuation that repeat the first table's leading
    rows: the header, plus any header-like row the detector left in the body
    ("At least | But less than | Your credit is–")."""
    n = 0
    while n < min(first.n_rows, cont.n_rows) and _row_key(first, n) == _row_key(cont, n) and _row_key(cont, n):
        n += 1
    return max(n, len(first.header_rows))


def _page(node) -> int | None:
    return node.source_regions[0]["page"] if node.source_regions else None


_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


def _sheet_title(wb: Workbook, base: str) -> str:
    base = _SHEET_BAD.sub(" ", base)[:31] or "Table"
    title, k = base, 2
    while title in wb.sheetnames:
        suffix = f" ({k})"
        title = base[:31 - len(suffix)] + suffix
        k += 1
    return title


def _write_table_group(wb: Workbook, group: list[TableNode], stats: dict) -> None:
    first = group[0]
    pages = sorted({p for t in group for p in [_page(t)] if p is not None})
    span = f"{pages[0]}" if len(pages) == 1 else f"{pages[0]}-{pages[-1]}"
    stats["_table_no"] = stats.get("_table_no", 0) + 1
    ws = wb.create_sheet(_sheet_title(wb, f"p{span} table {stats['_table_no']}"))
    hdr = len(first.header_rows)
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    offset = 0
    owner: dict[tuple[int, int], tuple[int, int]] = {}   # sheet cell -> anchor of the span that covers it
    overlaps = 0
    for k, t in enumerate(group):
        skip = _repeated_rows(first, t) if k > 0 else 0   # a continuation repeats the header: written once
        for c in sorted(t.cells, key=lambda c: (c["row"], c["col"])):
            if c["row"] < skip:
                continue
            r = c["row"] - skip + offset + 1
            col = c["col"] + 1
            span = [(rr, cc) for rr in range(r, r + c["rowspan"]) for cc in range(col, col + c["colspan"])]
            taken = [owner[p] for p in span if p in owner]
            if taken:
                # a cell inside an earlier span: its text joins the span's owner, never dropped
                ar, ac = taken[0]
                anchor = ws.cell(ar, ac)
                if c["raw_text"].strip():
                    anchor.value = _clean(f"{anchor.value}\n{c['raw_text']}" if anchor.value is not None else c["raw_text"])
                    anchor.data_type = "s"
                    anchor.alignment = wrap
                overlaps += 1
                continue
            for p in span:
                owner[p] = (r, col)
            cell = ws.cell(r, col)
            kind = _put(cell, c["raw_text"])
            stats["cells_number" if kind == "number" else "cells_text"] += 1
            if c["raw_text"] and "\n" in c["raw_text"]:
                cell.alignment = wrap
            if c["row"] in t.header_rows or c.get("bold"):
                cell.font = bold
                if kind == "text":
                    cell.alignment = wrap
            if c["rowspan"] > 1 or c["colspan"] > 1:
                ws.merge_cells(start_row=r, start_column=col,
                               end_row=r + c["rowspan"] - 1, end_column=col + c["colspan"] - 1)
                stats["merged_cells"] += 1
        offset += t.n_rows - skip
    # column widths from what the cells display in Excel's default font: every
    # number whole (a narrow column shows "####"), text at its longest word
    # (it wraps), capped so a long label does not make a page-wide column
    widths: dict[int, float] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or type(cell).__name__ == "MergedCell":
                continue
            if isinstance(cell.value, (int, float)):
                need = len(display(cell.value, cell.number_format)) + 2
            else:
                words = str(cell.value).split()
                need = min(max((len(w) for w in words), default=0) + 2, 40)
                if not any(cell.coordinate in r for r in ws.merged_cells.ranges) and "\n" not in str(cell.value):
                    need = min(len(str(cell.value)) + 2, 40)   # a single-line label reads whole
            widths[cell.column] = max(widths.get(cell.column, 4.0), need)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(60.0, w)
    if hdr and first.header_rows == list(range(hdr)):
        ws.freeze_panes = f"A{hdr + 1}"
    if overlaps:
        stats["warnings"].append({"code": "XLSX_CELL_OVERLAP", "sheet": ws.title, "count": overlaps,
                                  "detail": "cells inside another cell's span were merged into that cell's text"})
    stats["tables"] += 1
    stats["rows"] += offset
    if len(group) > 1:
        stats["warnings"].append({"code": "TABLE_JOINED", "tables": [t.id for t in group], "pages": pages,
                                  "detail": "tables with identical headers continuing across page halves or pages are one sheet; the repeated header is written once"})


def write_xlsx(doc, path: Path) -> dict:
    stats = {"tables": 0, "rows": 0, "cells_number": 0, "cells_text": 0, "merged_cells": 0,
             "text_rows": 0, "links": 0, "warnings": []}
    wb = Workbook()
    text_ws = wb.active
    text_ws.title = "Text"
    text_ws.append(["Page", "Kind", "Text"])
    for c in text_ws[1]:
        c.font = Font(bold=True)
    text_ws.freeze_panes = "A2"
    text_ws.column_dimensions["A"].width = 6
    text_ws.column_dimensions["B"].width = 12
    text_ws.column_dimensions["C"].width = 100
    groups: list[list[TableNode]] = []
    figures: dict[int, int] = {}
    multi_link: dict[int, int] = {}

    def text_row(page, kind, text, uri=None):
        r = text_ws.max_row + 1
        text_ws.cell(r, 1, page)
        text_ws.cell(r, 2, kind).data_type = "s"
        cell = text_ws.cell(r, 3)
        cell.value = _clean(text); cell.data_type = "s"
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if uri:
            cell.hyperlink = uri
            stats["links"] += 1
        stats["text_rows"] += 1

    for nid in doc.flow:
        node = doc.nodes[nid]
        page = _page(node)
        if isinstance(node, TableNode):
            if groups and _joinable(groups[-1][-1], node):
                groups[-1].append(node)
            else:
                groups.append([node])
        elif isinstance(node, TextNode):
            uris = [r.get("uri") for r in node.runs if r.get("uri")]
            if len(set(uris)) > 1:
                multi_link[page] = multi_link.get(page, 0) + len(set(uris)) - 1
            kind = "invisible text" if node.kind == "invisible" else node.kind
            text_row(page, kind, node.text, uris[0] if uris else None)
        elif isinstance(node, FormValueNode):
            if node.placement == "table_cell":
                continue                         # already inside its table cell
            shown = node.raw_value if node.checked is None else ("checked" if node.checked else "unchecked")
            text_row(page, "form value", f"{node.label or node.field_name}: {shown}")
        elif isinstance(node, CommentNode):
            who = f"{node.author}: " if node.author else ""
            text_row(page, "comment", f"{who}{node.text}")
        elif isinstance(node, FigureNode):
            figures[page] = figures.get(page, 0) + 1
            text_row(page, "picture", "[picture not carried into XLSX]" + (f" {node.absorbed_text}" if node.absorbed_text else ""))
    for g in groups:
        _write_table_group(wb, g, stats)
    if groups:
        wb.move_sheet("Text", offset=len(wb.sheetnames) - 1)   # tables first, the text sheet last
    else:
        stats["warnings"].append({"code": "XLSX_NO_TABLES", "detail": "no table was reconstructed; the text is on the Text sheet"})
    for page, n in sorted(figures.items()):
        stats["warnings"].append({"code": "XLSX_PICTURE_OMITTED", "page": page, "count": n,
                                  "detail": "pictures are not carried into XLSX; a placeholder row marks each on the Text sheet"})
    for page, n in sorted(multi_link.items()):
        stats["warnings"].append({"code": "XLSX_LINKS_REDUCED", "page": page, "count": n,
                                  "detail": "a cell holds one hyperlink; further links in the same paragraph stay plain text"})
    wb.active = 0
    wb.save(str(path))
    stats.pop("_table_no", None)
    return stats
